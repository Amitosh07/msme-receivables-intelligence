"""Historical CSV/XLSX payment validation and idempotent ingestion."""

from __future__ import annotations

import csv
import io
import logging
import math
import re
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Optional
from zipfile import BadZipFile

from fastapi import HTTPException, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, literal, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from backend.app.models.invoice import Invoice
from backend.app.models.payment import Payment
from backend.app.schemas.payment import (
    ImportErrorRow,
    PaymentImportResponse,
    PaymentImportRowResult,
)
from backend.app.services.customer_identity import (
    normalize_customer_name,
    normalize_gstin,
    resolve_customer_identity,
    resolve_or_create_historical_customer,
)

logger = logging.getLogger(__name__)


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "customer_name": (
        "customer_name", "customer", "company", "company_name", "client",
        "client_name", "buyer", "buyer_name", "party", "party_name", "name_customer",
    ),
    "invoice": (
        "invoice_number", "invoice_no", "invoice", "invoice_reference", "bill_number",
        "bill_no", "invoicenumber", "invoiceno",
    ),
    "date": (
        "payment_date", "paid_date", "received_date", "date_received", "paymentdate",
        "date_paid", "clear_date", "date",
    ),
    "amount": (
        "payment_amount", "paid_amount", "received_amount", "amount_paid",
        "amount_received", "paymentamount", "amount",
    ),
    "gstin": ("gstin", "gstin_no", "gst_no", "gst_number", "gstin_uin"),
    "customer_id": ("customer_id", "customer_uuid"),
    "customer_ref": (
        "customer_reference", "customer_ref", "customer_number", "customer_no", "cust_number",
    ),
    "payment_reference": (
        "reference", "payment_reference", "paymentreference", "utr", "transaction_id",
    ),
}

INDIA_DATE_FORMATS = (
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d %b %Y",
    "%d %B %Y",
)
UNAMBIGUOUS_TEXT_DATE_FORMATS = (
    "%b %d, %Y",
    "%B %d, %Y",
)
ISO_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y-%m-%d %H:%M:%S",
)
_PLAIN_NUMBER = re.compile(r"^\d+(?:\.\d+)?$")
_WESTERN_GROUPED_NUMBER = re.compile(r"^\d{1,3}(?:,\d{3})+(?:\.\d+)?$")
_INDIAN_GROUPED_NUMBER = re.compile(r"^\d{1,3}(?:,\d{2})*,\d{3}(?:\.\d+)?$")


@dataclass(frozen=True)
class NormalizedInputRow:
    row_number: int
    values: dict[str, Any]


@dataclass(frozen=True)
class PaymentInputTable:
    headers: list[str]
    rows: list[NormalizedInputRow]
    source_type: str


def _normalize_header(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "").strip()
    value = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", value)
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")


def _find_column(fieldnames: list[str], group: str) -> Optional[str]:
    normalized = {_normalize_header(field): field for field in fieldnames if field}
    for alias in HEADER_ALIASES[group]:
        if alias in normalized:
            return normalized[alias]
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _normalize_reference_for_invoice_match(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "").strip()
    return re.sub(r"\s+", "", value).casefold()


def _parse_date(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or timezone.utc)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc)
    if value is None or isinstance(value, (int, float, Decimal)):
        return None

    cleaned = unicodedata.normalize("NFKC", str(value)).strip()
    cleaned = re.sub(r"(?<=\d)(st|nd|rd|th)\b", "", cleaned, flags=re.I)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T].*", cleaned):
        try:
            parsed = datetime.fromisoformat(cleaned.replace("Z", "+00:00"))
            return parsed.replace(tzinfo=parsed.tzinfo or timezone.utc)
        except ValueError:
            return None
    for fmt in INDIA_DATE_FORMATS + ISO_DATE_FORMATS + UNAMBIGUOUS_TEXT_DATE_FORMATS:
        try:
            parsed = datetime.strptime(cleaned, fmt)
            return parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _parse_amount(value: Any) -> Optional[Decimal]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
        except (InvalidOperation, ValueError):
            return None
    else:
        cleaned = unicodedata.normalize("NFKC", str(value)).strip()
        negative_parentheses = cleaned.startswith("(") and cleaned.endswith(")")
        if negative_parentheses:
            cleaned = cleaned[1:-1].strip()
        cleaned = re.sub(r"(?i)^\s*(?:INR|USD|CAD|EUR|GBP|AUD|RS\.?)\s*", "", cleaned)
        cleaned = re.sub(r"(?i)\s*(?:INR|USD|CAD|EUR|GBP|AUD|RS\.?)\s*$", "", cleaned)
        cleaned = re.sub(r"[₹$€£\s]", "", cleaned)
        sign = ""
        if cleaned[:1] in ("+", "-"):
            sign, cleaned = cleaned[0], cleaned[1:]
        if not (
            _PLAIN_NUMBER.fullmatch(cleaned)
            or _WESTERN_GROUPED_NUMBER.fullmatch(cleaned)
            or _INDIAN_GROUPED_NUMBER.fullmatch(cleaned)
        ):
            return None
        cleaned = sign + cleaned.replace(",", "")
        if negative_parentheses:
            cleaned = "-" + cleaned.lstrip("+")
        try:
            amount = Decimal(cleaned)
        except (InvalidOperation, ValueError):
            return None

    if not amount.is_finite() or amount <= 0:
        return None
    try:
        return amount.quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _decode_csv(file_content: bytes) -> str:
    if not file_content:
        raise HTTPException(status_code=400, detail="Uploaded CSV file is empty.")
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return file_content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=400, detail="Unable to decode CSV. Save it as UTF-8 and try again.")


def _read_csv(file_content: bytes) -> PaymentInputTable:
    decoded = _decode_csv(file_content)
    header_line = next((line for line in decoded.splitlines() if line.strip()), "")
    delimiter = max((",", ";", "\t", "|"), key=header_line.count)
    if header_line.count(delimiter) == 0:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(decoded, newline=""), delimiter=delimiter)
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="Uploaded CSV file contains no headers.")
    headers = [str(header or "") for header in reader.fieldnames]
    rows = [
        NormalizedInputRow(row_number=row_number, values=dict(row))
        for row_number, row in enumerate(reader, start=2)
        if any(_cell_text(value) for value in row.values())
    ]
    return PaymentInputTable(headers=headers, rows=rows, source_type="CSV")


def _read_xlsx(file_content: bytes) -> PaymentInputTable:
    if not file_content:
        raise HTTPException(status_code=400, detail="Uploaded XLSX file is empty.")
    try:
        workbook = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Unable to read XLSX workbook.") from exc
    try:
        worksheet = workbook.active
        header_values: tuple[Any, ...] | None = None
        header_row_number = 0
        iterator = worksheet.iter_rows(values_only=True)
        for row_number, values in enumerate(iterator, start=1):
            if any(_cell_text(value) for value in values):
                header_values = values
                header_row_number = row_number
                break
        if header_values is None:
            raise HTTPException(status_code=400, detail="Uploaded XLSX workbook contains no headers.")
        headers = [_cell_text(value) for value in header_values]
        rows: list[NormalizedInputRow] = []
        for row_number, values in enumerate(iterator, start=header_row_number + 1):
            if not any(_cell_text(value) for value in values):
                continue
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            rows.append(NormalizedInputRow(row_number=row_number, values=row))
        return PaymentInputTable(headers=headers, rows=rows, source_type="XLSX")
    finally:
        workbook.close()


def _normalize_index_text(
    db: Session,
    value: str | None,
    cache: dict[str, str] | None = None,
) -> str:
    """Use the live PostgreSQL expression that defines invoice-reference uniqueness."""
    raw_value = value or ""
    if cache is not None and raw_value in cache:
        return cache[raw_value]
    normalized = db.scalar(
        select(func.lower(func.btrim(func.coalesce(literal(raw_value), literal("")))))
    ) or ""
    if cache is not None:
        cache[raw_value] = normalized
    return normalized


def _payment_key(
    customer_identity_key: str,
    normalized_reference: str,
    payment_date: datetime,
    amount: Decimal | float,
) -> tuple[str, str, datetime, Decimal]:
    return (
        customer_identity_key,
        normalized_reference,
        payment_date.astimezone(timezone.utc),
        Decimal(str(amount)).quantize(Decimal("0.01")),
    )


def _parse_customer_id(value: Any) -> uuid.UUID | None:
    text_value = _cell_text(value)
    if not text_value:
        return None
    try:
        return uuid.UUID(text_value)
    except ValueError as exc:
        raise ValueError("Cannot import: customer_id is not a valid UUID.") from exc


def _identity_label(
    *,
    customer_id: uuid.UUID | None,
    gstin: str,
    customer_ref: str,
    customer_name: str,
) -> str:
    if customer_id:
        return str(customer_id)
    valid_gstin = normalize_gstin(gstin)
    if valid_gstin:
        return valid_gstin
    return customer_ref or customer_name or "unresolved customer"


def _reject_row(
    *,
    row: NormalizedInputRow,
    reason: str,
    errors: list[ImportErrorRow],
    row_results: list[PaymentImportRowResult],
) -> None:
    errors.append(ImportErrorRow(row_number=row.row_number, reason=reason, raw_data=row.values))
    row_results.append(PaymentImportRowResult(
        row_number=row.row_number,
        status="rejected",
        reason=reason,
        raw_data=row.values,
    ))


def _process_table(
    db: Session,
    business_id: uuid.UUID,
    table: PaymentInputTable,
    *,
    persist: bool,
    customer_id: Optional[uuid.UUID] = None,
) -> PaymentImportResponse:
    inv_col = _find_column(table.headers, "invoice")
    date_col = _find_column(table.headers, "date")
    amount_col = _find_column(table.headers, "amount")
    payment_ref_col = _find_column(table.headers, "payment_reference")
    customer_ref_col = _find_column(table.headers, "customer_ref")
    customer_name_col = _find_column(table.headers, "customer_name")
    gstin_col = _find_column(table.headers, "gstin")
    customer_id_col = _find_column(table.headers, "customer_id")

    missing = [
        canonical
        for canonical, found in (("payment_date", date_col), ("payment_amount", amount_col))
        if not found
    ]
    if customer_id is None and not any((customer_name_col, gstin_col, customer_id_col, customer_ref_col)):
        missing.append("customer identity")
    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"{table.source_type} is missing required columns: {', '.join(missing)}.",
        )

    invoices = list(db.scalars(select(Invoice).where(Invoice.business_id == business_id)).all())
    invoices_by_ref = {
        _normalize_reference_for_invoice_match(invoice.invoice_number): invoice
        for invoice in invoices
    }
    invoice_number_by_id = {invoice.id: invoice.invoice_number for invoice in invoices}

    normalization_cache: dict[str, str] = {}
    existing_keys = {
        _payment_key(
            payment.customer_identity_key,
            _normalize_index_text(db, payment.invoice_reference, normalization_cache),
            payment.payment_date,
            payment.amount,
        )
        for payment in db.scalars(select(Payment).where(Payment.business_id == business_id)).all()
        if payment.customer_identity_key is not None
    }

    preview_scope = db.begin_nested() if not persist else None
    file_keys: set[tuple[str, str, datetime, Decimal]] = set()
    invoices_to_mark_paid: set[uuid.UUID] = set()
    errors: list[ImportErrorRow] = []
    row_results: list[PaymentImportRowResult] = []
    imported = rejected = duplicate_file = duplicate_existing = unmatched = matched = 0

    try:
        for row in table.rows:
            values = row.values
            raw_invoice = _cell_text(values.get(inv_col)) if inv_col else ""
            raw_date = values.get(date_col) if date_col else None
            raw_amount = values.get(amount_col) if amount_col else None
            customer_name = _cell_text(values.get(customer_name_col)) if customer_name_col else ""
            customer_ref = _cell_text(values.get(customer_ref_col)) if customer_ref_col else ""
            gstin = _cell_text(values.get(gstin_col)) if gstin_col else ""

            parsed_date = _parse_date(raw_date)
            if parsed_date is None:
                rejected += 1
                _reject_row(
                    row=row,
                    reason=f"Cannot import: invalid or unresolvable payment date '{_cell_text(raw_date)}'.",
                    errors=errors,
                    row_results=row_results,
                )
                continue
            parsed_amount = _parse_amount(raw_amount)
            if parsed_amount is None:
                rejected += 1
                _reject_row(
                    row=row,
                    reason="Cannot import: payment amount must be finite and greater than zero.",
                    errors=errors,
                    row_results=row_results,
                )
                continue

            try:
                parsed_cust_id = _parse_customer_id(
                    values.get(customer_id_col) if customer_id_col else None
                )
                requested_customer_id = parsed_cust_id if parsed_cust_id is not None else customer_id
            except ValueError as exc:
                rejected += 1
                _reject_row(row=row, reason=str(exc), errors=errors, row_results=row_results)
                continue

            matching_invoice = (
                invoices_by_ref.get(_normalize_reference_for_invoice_match(raw_invoice))
                if raw_invoice
                else None
            )
            identity_result = resolve_customer_identity(
                db,
                business_id=business_id,
                customer_id=requested_customer_id,
                gstin=gstin or None,
                customer_ref=customer_ref or None,
                display_name=customer_name or None,
            )
            resolved_customer = identity_result.customer

            has_identity_evidence = bool(
                requested_customer_id
                or normalize_gstin(gstin)
                or customer_ref
                or normalize_customer_name(customer_name)
            )
            if matching_invoice is not None and matching_invoice.customer_id is not None:
                invoice_customer = matching_invoice.customer
                if resolved_customer is None and has_identity_evidence:
                    row_label = _identity_label(
                        customer_id=requested_customer_id,
                        gstin=gstin,
                        customer_ref=customer_ref,
                        customer_name=customer_name,
                    )
                    rejected += 1
                    _reject_row(
                        row=row,
                        reason=(
                            "Cannot import: payment customer identity conflicts with existing "
                            f"invoice's assigned customer. Row indicates '{row_label}', invoice is "
                            f"linked to '{invoice_customer.display_name}'. Manual review required."
                        ),
                        errors=errors,
                        row_results=row_results,
                    )
                    continue
                if resolved_customer is not None and resolved_customer.id != matching_invoice.customer_id:
                    rejected += 1
                    _reject_row(
                        row=row,
                        reason=(
                            "Cannot import: payment customer identity conflicts with existing "
                            "invoice's assigned customer. Row indicates "
                            f"'{resolved_customer.display_name}', invoice is linked to "
                            f"'{invoice_customer.display_name}'. Manual review required."
                        ),
                        errors=errors,
                        row_results=row_results,
                    )
                    continue

            if resolved_customer is None:
                try:
                    identity_result = resolve_or_create_historical_customer(
                        db,
                        business_id=business_id,
                        customer_id=requested_customer_id,
                        gstin=gstin or None,
                        customer_ref=customer_ref or None,
                        display_name=customer_name or None,
                    )
                    resolved_customer = identity_result.customer
                except ValueError as exc:
                    rejected += 1
                    _reject_row(row=row, reason=str(exc), errors=errors, row_results=row_results)
                    continue

            if resolved_customer is None:
                rejected += 1
                _reject_row(
                    row=row,
                    reason="Cannot import: no resolvable customer identity found for this row.",
                    errors=errors,
                    row_results=row_results,
                )
                continue

            customer_identity_key = f"customer:{resolved_customer.id}"
            stored_reference = (
                matching_invoice.invoice_number
                if matching_invoice is not None
                else (raw_invoice or None)
            )
            normalized_invoice_reference = _normalize_index_text(
                db, stored_reference, normalization_cache
            )
            key = _payment_key(
                customer_identity_key,
                normalized_invoice_reference,
                parsed_date,
                parsed_amount,
            )
            if key in file_keys:
                duplicate_file += 1
                row_results.append(PaymentImportRowResult(
                    row_number=row.row_number,
                    status="duplicate",
                    reason="Duplicate payment within the uploaded file.",
                    payment_date=parsed_date,
                    invoice_number=matching_invoice.invoice_number if matching_invoice else (raw_invoice or None),
                    invoice_id=matching_invoice.id if matching_invoice else None,
                    raw_data=values,
                ))
                continue
            file_keys.add(key)
            if key in existing_keys:
                duplicate_existing += 1
                row_results.append(PaymentImportRowResult(
                    row_number=row.row_number,
                    status="duplicate",
                    reason="Duplicate payment already exists.",
                    payment_date=parsed_date,
                    invoice_number=matching_invoice.invoice_number if matching_invoice else (raw_invoice or None),
                    invoice_id=matching_invoice.id if matching_invoice else None,
                    raw_data=values,
                ))
                continue

            payment = Payment(
                business_id=business_id,
                invoice_id=matching_invoice.id if matching_invoice is not None else None,
                invoice_reference=stored_reference,
                payment_date=parsed_date,
                amount=parsed_amount,
                reference=(
                    _cell_text(values.get(payment_ref_col)) or None if payment_ref_col else None
                ),
                customer_identity_key=customer_identity_key,
                provenance="import",
            )
            try:
                with db.begin_nested():
                    db.add(payment)
                    db.flush()
            except IntegrityError as exc:
                diagnostic = getattr(getattr(exc, "orig", None), "diag", None)
                if getattr(diagnostic, "constraint_name", None) != "uq_payment_natural_key":
                    db.rollback()
                    logger.exception("Failed to persist payment-history import row %d", row.row_number)
                    raise HTTPException(
                        status_code=500,
                        detail="Failed to persist payment records to the database.",
                    ) from exc
                duplicate_existing += 1
                row_results.append(PaymentImportRowResult(
                    row_number=row.row_number,
                    status="duplicate",
                    reason="Duplicate payment conflicted with an existing database record.",
                    payment_date=parsed_date,
                    invoice_number=matching_invoice.invoice_number if matching_invoice else (raw_invoice or None),
                    invoice_id=matching_invoice.id if matching_invoice else None,
                    raw_data=values,
                ))
                continue

            imported += 1
            is_unmatched = matching_invoice is None
            if is_unmatched:
                unmatched += 1
                result_reason = (
                    "Payment imported; invoice reference did not match an existing invoice."
                    if raw_invoice
                    else "Payment imported without an invoice reference; customer history retained."
                )
            else:
                matched += 1
                invoices_to_mark_paid.add(matching_invoice.id)
                result_reason = "Imported historical payment."
            row_results.append(PaymentImportRowResult(
                row_number=row.row_number,
                status="imported",
                reason=result_reason if persist else f"Valid for import. {result_reason}",
                unmatched=is_unmatched,
                payment_id=payment.id,
                payment_date=parsed_date,
                invoice_number=matching_invoice.invoice_number if matching_invoice else (raw_invoice or None),
                invoice_id=matching_invoice.id if matching_invoice else None,
                raw_data=values,
            ))

        if persist:
            try:
                for invoice_id in invoices_to_mark_paid:
                    invoices_by_ref[
                        _normalize_reference_for_invoice_match(invoice_number_by_id[invoice_id])
                    ].payment_status = "PAID"
                db.commit()
            except Exception as exc:
                db.rollback()
                logger.exception("Failed to persist payment-history import")
                raise HTTPException(
                    status_code=500,
                    detail="Failed to persist payment records to the database.",
                ) from exc

        duplicates = duplicate_file + duplicate_existing
        response = PaymentImportResponse(
            total_rows=len(table.rows),
            valid_rows=len(table.rows) - rejected,
            imported=imported,
            matched=matched,
            duplicates=duplicates,
            duplicates_in_file=duplicate_file,
            duplicates_existing=duplicate_existing,
            rejected=rejected,
            unmatched=unmatched,
            preview=not persist,
            errors=errors[:50],
            row_results=row_results,
        )
        logger.info(
            "Payment %s %s: total=%d imported=%d matched=%d unmatched=%d "
            "file_duplicates=%d database_duplicates=%d rejected=%d",
            table.source_type,
            "import" if persist else "preview",
            len(table.rows),
            imported,
            matched,
            unmatched,
            duplicate_file,
            duplicate_existing,
            rejected,
        )
        return response
    finally:
        if preview_scope is not None and preview_scope.is_active:
            preview_scope.rollback()


def preview_payments_file(
    db: Session,
    business_id: uuid.UUID,
    file_content: bytes,
    file_type: str,
    customer_id: Optional[uuid.UUID] = None,
) -> PaymentImportResponse:
    table = _read_xlsx(file_content) if file_type == "xlsx" else _read_csv(file_content)
    return _process_table(db, business_id, table, persist=False, customer_id=customer_id)


def import_payments_file(
    db: Session,
    business_id: uuid.UUID,
    file_content: bytes,
    file_type: str,
    customer_id: Optional[uuid.UUID] = None,
) -> PaymentImportResponse:
    table = _read_xlsx(file_content) if file_type == "xlsx" else _read_csv(file_content)
    return _process_table(db, business_id, table, persist=True, customer_id=customer_id)


def preview_payments_csv(
    db: Session,
    business_id: uuid.UUID,
    file_content: bytes,
) -> PaymentImportResponse:
    return preview_payments_file(db, business_id, file_content, "csv")


def import_payments_csv(
    db: Session,
    business_id: uuid.UUID,
    file_content: bytes,
) -> PaymentImportResponse:
    return import_payments_file(db, business_id, file_content, "csv")
